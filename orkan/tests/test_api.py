"""Web-Oberfläche: Upload, Abgleich, Downloads, Queue, Admin, Abschottung des Pools."""
from __future__ import annotations

import io
import time
from pathlib import Path

import pandas as pd
import pytest
from openpyxl import load_workbook

from app import config, queue_store

ADMIN = ("pruefer", "geheim")


def _csv(inhalt: str) -> tuple[str, io.BytesIO, str]:
    return ("liste.csv", io.BytesIO(inhalt.encode("utf-8")), "text/csv")


def _excel(df: pd.DataFrame) -> tuple[str, io.BytesIO, str]:
    puffer = io.BytesIO()
    df.to_excel(puffer, index=False)
    puffer.seek(0)
    return ("liste.xlsx", puffer, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


def _durchlauf(klient, datei, zuordnung=None, kunde="Testbetrieb"):
    hoch = klient.post("/api/upload", files={"datei": datei})
    assert hoch.status_code == 200, hoch.text
    job = hoch.json()
    antwort = klient.post("/api/verarbeiten", json={
        "job_id": job["job_id"],
        "zuordnung": zuordnung or job["zuordnung"],
        "kundenkennung": kunde,
    })
    assert antwort.status_code == 200, antwort.text
    return job, antwort.json()


# ---------------------------------------------------------------------------
# Abnahmekriterium 1: CSV mit "Artikelnummer", 1.000 Zeilen, unter 10 Sekunden
# ---------------------------------------------------------------------------
def test_tausend_zeilen_in_unter_zehn_sekunden(klient):
    nummern = ["CUS18H", "CUS22H", "CUS28H", "GWS-M10", "01.001.90"] * 200
    inhalt = "Artikelnummer;Bezeichnung\n" + "\n".join(f"{n};Testartikel" for n in nummern)

    start = time.perf_counter()
    _, ergebnis = _durchlauf(klient, _csv(inhalt))
    dauer = time.perf_counter() - start

    assert ergebnis["zusammenfassung"]["gesamt"] == 1000
    assert ergebnis["zusammenfassung"]["gefunden"] == 1000
    assert ergebnis["zusammenfassung"]["trefferquote"] == 100
    assert dauer < 10, f"Verarbeitung dauerte {dauer:.1f}s"


# ---------------------------------------------------------------------------
# Abnahmekriterium 2: Excel mit abweichenden Headern
# ---------------------------------------------------------------------------
def test_excel_mit_abweichenden_headern(klient):
    df = pd.DataFrame({
        "Art-Nr.": ["CUS18H", "GWS-M10"],
        "Lieferant": ["Mein Großhandel", "Adolf Würth"],
        "Bezeichnung": ["Bogen 90", "Gewindestange"],
    })
    hoch = klient.post("/api/upload", files={"datei": _excel(df)})
    assert hoch.status_code == 200
    zuordnung = hoch.json()["zuordnung"]

    assert zuordnung["Artikelnummer"] == "Art-Nr."
    assert zuordnung["Lieferant"] == "Lieferant"
    assert zuordnung["Bezeichnung1"] == "Bezeichnung"


def test_nutzer_kann_die_zuordnung_ueberschreiben(klient):
    inhalt = "Spalte A;Spalte B\nBogen 90;CUS18H\nMuffe;CUS28H"
    hoch = klient.post("/api/upload", files={"datei": _csv(inhalt)})
    job = hoch.json()

    antwort = klient.post("/api/verarbeiten", json={
        "job_id": job["job_id"],
        "zuordnung": {"Artikelnummer": "Spalte B", "Bezeichnung1": "Spalte A"},
    })
    assert antwort.json()["zusammenfassung"]["gefunden"] == 2


# ---------------------------------------------------------------------------
# Abnahmekriterium 3: Unbekannte Nummern → Importdatei UND Queue
# ---------------------------------------------------------------------------
def test_unbekannte_nummern_landen_in_importdatei_und_queue(klient):
    inhalt = (
        "Artikelnummer;Bezeichnung;Artikelgruppe;Lieferant\n"
        "CUS18H;Bogen 90;Pressfittings;Sprügel\n"
        "NEU-KPF-2815;Kupfer Pressfitting Bogen 90;Pressfittings;Sprügel\n"
    )
    _, ergebnis = _durchlauf(klient, _csv(inhalt), kunde="Mustermann")

    z = ergebnis["zusammenfassung"]
    assert z["gefunden"] == 1 and z["neu"] == 1 and z["zur_pruefung"] == 1
    assert z["in_importdatei"] == 2

    # in der Importdatei
    antwort = klient.get(ergebnis["downloads"]["csv"])
    assert antwort.status_code == 200
    tabelle = pd.read_csv(io.BytesIO(antwort.content), sep=";", dtype=str, encoding="utf-8-sig")
    assert set(tabelle["Artikelnummer des Lieferanten"]) == {"CUS18H", "NEU-KPF-2815"}

    neu = tabelle[tabelle["Artikelnummer des Lieferanten"] == "NEU-KPF-2815"].iloc[0]
    assert neu["Lagerort"] == "Kupferpress"
    assert int(neu["Lager Min"]) >= 1 and int(neu["Lager Max"]) > int(neu["Lager Min"])
    assert pd.isna(neu["Artikelbild"])

    # in der Queue, mit Zeitstempel und Kundenkennung
    queue = queue_store.lesen()
    assert len(queue) == 1
    eintrag = queue.iloc[0]
    assert eintrag["ArtikelNrLieferant"] == "NEU-KPF-2815"
    assert eintrag["Kundenkennung"] == "Mustermann"
    assert eintrag["QueueStatus"] == queue_store.STATUS_OFFEN
    assert eintrag["Zeitstempel"]


def test_pool_wird_durch_den_abgleich_nicht_veraendert(klient):
    vorher = Path(config.POOL_PFAD).read_bytes()
    _durchlauf(klient, _csv("Artikelnummer\nNEU-IRGENDWAS-1\nNEU-IRGENDWAS-2"))
    assert Path(config.POOL_PFAD).read_bytes() == vorher


# ---------------------------------------------------------------------------
# Abnahmekriterium 4: Excel-Datei öffnet sauber
# ---------------------------------------------------------------------------
def test_excel_download_haelt_umlaute_und_lange_nummern(klient):
    _, ergebnis = _durchlauf(klient, _csv("Artikelnummer\n387500140020"))

    antwort = klient.get(ergebnis["downloads"]["xlsx"])
    assert antwort.status_code == 200
    blatt = load_workbook(io.BytesIO(antwort.content))["Import"]
    kopf = [z.value for z in blatt[1]]

    assert len(kopf) == 39
    assert blatt.cell(row=2, column=kopf.index("Bezeichnung1") + 1).value == "Tauchsägeblatt"
    ean = blatt.cell(row=2, column=kopf.index("EAN") + 1)
    assert ean.value == "4004478123456" and ean.number_format == "@"


# ---------------------------------------------------------------------------
# Abnahmekriterium 5: Der Pool ist über keine Route abrufbar
# ---------------------------------------------------------------------------
@pytest.mark.parametrize(
    "pfad",
    [
        "/daten/Datenpool_bereinigt.csv", "/pool", "/api/pool", "/api/pool/suche",
        "/api/artikel", "/static/../daten/Datenpool_bereinigt.csv", "/openapi.json",
        "/docs", "/api/download/x/pool",
    ],
)
def test_pool_ist_ueber_keine_route_abrufbar(klient, pfad):
    antwort = klient.get(pfad)
    assert antwort.status_code >= 400, f"{pfad} lieferte {antwort.status_code}"


def test_status_verraet_nur_die_anzahl_keine_artikel(klient):
    daten = klient.get("/api/status").json()
    assert daten["artikel_im_pool"] > 0
    text = str(daten)
    assert "CUS18H" not in text and "Sprügel" not in text


def test_ergebnis_enthaelt_nur_die_eigenen_nummern(klient):
    _, ergebnis = _durchlauf(klient, _csv("Artikelnummer\nCUS18H"))
    nummern = {z["Artikelnummer"] for z in ergebnis["report"]}
    assert nummern == {"CUS18H"}
    assert "GWS-M10" not in str(ergebnis)


def test_zu_viele_zeilen_werden_abgelehnt(klient, monkeypatch):
    monkeypatch.setattr(config, "MAX_ZEILEN", 5)
    inhalt = "Artikelnummer\n" + "\n".join("CUS18H" for _ in range(10))
    antwort = klient.post("/api/upload", files={"datei": _csv(inhalt)})
    assert antwort.status_code == 400
    assert "Zeilen" in antwort.json()["fehler"]


# ---------------------------------------------------------------------------
# Uploads werden gelöscht
# ---------------------------------------------------------------------------
def test_hochgeladene_datei_wird_nach_dem_einlesen_geloescht(klient):
    job, _ = _durchlauf(klient, _csv("Artikelnummer\nCUS18H"))
    verzeichnis = Path(config.ARBEITSVERZEICHNIS) / job["job_id"]
    uebrig = {p.name for p in verzeichnis.iterdir()}

    assert not any(n.startswith("upload") for n in uebrig)
    assert "kundenliste.pkl" not in uebrig
    assert any(n.startswith("Importdatei_") for n in uebrig)


# ---------------------------------------------------------------------------
# Fehlerfälle
# ---------------------------------------------------------------------------
def test_leere_datei_meldet_klaren_fehler(klient):
    antwort = klient.post("/api/upload", files={"datei": _csv("")})
    assert antwort.status_code == 400
    assert "leer" in antwort.json()["fehler"].lower()


def test_unbekanntes_format_wird_abgelehnt(klient):
    datei = ("liste.docx", io.BytesIO(b"abc"), "application/octet-stream")
    antwort = klient.post("/api/upload", files={"datei": datei})
    assert antwort.status_code == 400
    assert "Excel" in antwort.json()["fehler"]


def test_unbekannter_auftrag_liefert_404(klient):
    antwort = klient.post("/api/verarbeiten", json={"job_id": "gibtesnicht", "zuordnung": {}})
    assert antwort.status_code == 404


def test_ohne_artikelnummer_spalte_kommt_ein_hinweis(klient):
    hoch = klient.post("/api/upload", files={"datei": _csv("A;B\nx;y\np;q")})
    antwort = klient.post("/api/verarbeiten", json={
        "job_id": hoch.json()["job_id"], "zuordnung": {"Bezeichnung1": "A"},
    })
    assert antwort.status_code == 400
    assert "Artikelnummern" in antwort.json()["fehler"]


# ---------------------------------------------------------------------------
# Admin
# ---------------------------------------------------------------------------
def test_admin_braucht_anmeldung(klient):
    assert klient.get("/admin/neuaufnahmen").status_code == 401
    assert klient.post("/admin/neuaufnahmen/uebernehmen", json={"queue_ids": []}).status_code == 401
    assert klient.post("/admin/pool/neuladen").status_code == 401


def test_admin_lehnt_falsches_passwort_ab(klient):
    assert klient.get("/admin/neuaufnahmen", auth=("pruefer", "falsch")).status_code == 401


def test_admin_zeigt_offene_neuaufnahmen(klient):
    _durchlauf(klient, _csv("Artikelnummer;Bezeichnung\nNEU-ADMIN-1;Kupfer Pressfitting"))
    seite = klient.get("/admin/neuaufnahmen", auth=ADMIN)

    assert seite.status_code == 200
    assert "NEU-ADMIN-1" in seite.text
    assert "Kupfer Pressfitting" in seite.text


def test_admin_uebernimmt_artikel_in_den_pool(klient):
    _durchlauf(klient, _csv("Artikelnummer;Bezeichnung\nNEU-ADMIN-2;Kupfer Pressfitting"))
    queue_id = queue_store.lesen(nur_offen=True).iloc[0]["QueueId"]

    vorher = klient.get("/api/status").json()["artikel_im_pool"]
    antwort = klient.post(
        "/admin/neuaufnahmen/uebernehmen", json={"queue_ids": [queue_id]}, auth=ADMIN
    )
    assert antwort.status_code == 200
    assert antwort.json()["uebernommen"] == 1
    assert antwort.json()["artikel_im_pool"] == vorher + 1

    # jetzt wird der Artikel direkt gefunden
    _, ergebnis = _durchlauf(klient, _csv("Artikelnummer\nNEU-ADMIN-2"))
    assert ergebnis["zusammenfassung"]["gefunden"] == 1
    assert queue_store.lesen(nur_offen=True).empty


def test_admin_kann_neuaufnahmen_verwerfen(klient):
    _durchlauf(klient, _csv("Artikelnummer\nNEU-ADMIN-3"))
    queue_id = queue_store.lesen(nur_offen=True).iloc[0]["QueueId"]

    vorher = klient.get("/api/status").json()["artikel_im_pool"]
    antwort = klient.post(
        "/admin/neuaufnahmen/verwerfen", json={"queue_ids": [queue_id]}, auth=ADMIN
    )
    assert antwort.json()["verworfen"] == 1
    assert klient.get("/api/status").json()["artikel_im_pool"] == vorher
    assert queue_store.lesen(nur_offen=True).empty


def test_admin_gibt_den_pool_nicht_heraus(klient):
    """Die Queue-Seite zeigt nur geprüfte Kandidaten, nie den Pool selbst."""
    _durchlauf(klient, _csv("Artikelnummer\nNEU-ADMIN-4"))
    seite = klient.get("/admin/neuaufnahmen", auth=ADMIN).text
    assert "NEU-ADMIN-4" in seite
    assert "CUS22H" not in seite and "01.001.90" not in seite


def test_pool_neuladen_nach_dateiaustausch(klient):
    assert klient.post("/admin/pool/neuladen", auth=ADMIN).json()["artikel_im_pool"] > 0


# ---------------------------------------------------------------------------
# Verschwundene Uploaddatei
# ---------------------------------------------------------------------------
# Die Ergebnisdateien werden nach ERGEBNIS_TTL_STUNDEN gelöscht. Steht der
# Kunde in dieser Zeit noch bei Schritt 2 und klickt dann auf „abgleichen",
# ist seine hochgeladene Liste weg. Früher gab es dafür einen nackten
# Serverfehler – obwohl die Meldung längst für ihn formuliert war.
def test_verschwundene_uploaddatei_ergibt_eine_lesbare_meldung(klient):
    antwort = klient.post("/api/upload", files={"datei": _csv("Artikelnummer\nCUS18H\n")})
    job_id = antwort.json()["job_id"]

    from app import jobs
    jobs.rohdaten_loeschen(jobs.hole(job_id))

    ergebnis = klient.post("/api/verarbeiten", json={"job_id": job_id})
    assert ergebnis.status_code == 410
    text = ergebnis.json()["fehler"]
    assert "erneut hoch" in text
    assert "Traceback" not in text


def test_die_meldung_kommt_im_gleichen_format_wie_andere_fehler(klient):
    """Die Oberfläche liest immer das Feld `fehler` – auch hier."""
    antwort = klient.post("/api/upload", files={"datei": _csv("Artikelnummer\nCUS18H\n")})
    job_id = antwort.json()["job_id"]
    from app import jobs
    jobs.rohdaten_loeschen(jobs.hole(job_id))

    ergebnis = klient.post("/api/verarbeiten", json={"job_id": job_id})
    assert set(ergebnis.json()) == {"fehler"}
