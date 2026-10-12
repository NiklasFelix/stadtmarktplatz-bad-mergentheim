"""Ausgabeformat: exakt 39 Spalten, feste Werte, Excel-Sicherheit."""
from __future__ import annotations

import csv

import pandas as pd
import pytest
from openpyxl import load_workbook

from app import config, export, matching

VORLAGE = config.VORLAGE_PFAD


def test_vorlage_hat_39_spalten():
    assert len(export.vorlagenspalten()) == 39


def test_vorlagenspalten_stimmen_mit_der_datei_ueberein():
    with VORLAGE.open("r", encoding="utf-8-sig", newline="") as f:
        aus_datei = [s.strip() for s in next(csv.reader(f, delimiter=";"))]
    assert export.vorlagenspalten() == aus_datei


def test_importtabelle_hat_exakt_die_vorlagenspalten(pool):
    ergebnis = matching.abgleich(pd.DataFrame({"Artikelnummer": ["CUS18H"]}), pool)
    tabelle = export.importtabelle_bauen(ergebnis.importzeilen)
    assert list(tabelle.columns) == export.vorlagenspalten()


def test_leere_liste_ergibt_gueltige_leere_tabelle():
    tabelle = export.importtabelle_bauen([])
    assert list(tabelle.columns) == export.vorlagenspalten()
    assert len(tabelle) == 0


def test_falsche_reihenfolge_wird_abgelehnt():
    spalten = export.vorlagenspalten()
    vertauscht = [spalten[1], spalten[0]] + spalten[2:]
    with pytest.raises(ValueError, match="Reihenfolge"):
        export.validiere_importformat(pd.DataFrame(columns=vertauscht))


def test_fehlende_spalte_wird_abgelehnt():
    with pytest.raises(ValueError, match="fehlende Spalten"):
        export.validiere_importformat(pd.DataFrame(columns=export.vorlagenspalten()[:-1]))


def test_unbekannte_spalte_wird_abgelehnt():
    with pytest.raises(ValueError, match="unbekannte Spalten"):
        export.validiere_importformat(
            pd.DataFrame(columns=export.vorlagenspalten() + ["Extra"])
        )


def test_feste_werte_und_leerfelder(pool):
    ergebnis = matching.abgleich(pd.DataFrame({"Artikelnummer": ["CUS18H"]}), pool)
    zeile = ergebnis.importzeilen[0]

    assert zeile["Satzart"] == "ST"
    assert zeile["Bestellfähig"] == "True"
    assert zeile["Bestand"] == "0"
    assert zeile["Lager/Gebäude"] == "Standard Lager"
    for leer in ("Artikelnr.", "EK-Preis", "Gewicht (Tara) in kg", "Höhe", "Breite", "Länge"):
        assert pd.isna(zeile[leer]), leer


# ---------------------------------------------------------------------------
# Dateien
# ---------------------------------------------------------------------------
def test_csv_ist_semikolon_getrennt_mit_bom(tmp_path, pool):
    ergebnis = matching.abgleich(pd.DataFrame({"Artikelnummer": ["CUS18H"]}), pool)
    tabelle = export.importtabelle_bauen(ergebnis.importzeilen)
    dateien = export.schreibe_importdatei(tabelle, tmp_path, "Test")

    roh = dateien["csv"].read_bytes()
    assert roh.startswith(b"\xef\xbb\xbf")  # UTF-8-BOM, damit Excel Umlaute erkennt
    text = roh.decode("utf-8-sig")
    assert text.splitlines()[0].split(";")[:2] == ["Artikelnr.", "Bezeichnung1"]


def test_csv_uebersteht_umlaute_und_lange_nummern(tmp_path, pool):
    ergebnis = matching.abgleich(pd.DataFrame({"Artikelnummer": ["387500140020"]}), pool)
    tabelle = export.importtabelle_bauen(ergebnis.importzeilen)
    dateien = export.schreibe_importdatei(tabelle, tmp_path, "Test")

    zurueck = pd.read_csv(dateien["csv"], sep=";", dtype=str, encoding="utf-8-sig")
    assert zurueck["Bezeichnung1"].iloc[0] == "Tauchsägeblatt"
    assert zurueck["Artikelnummer des Lieferanten"].iloc[0] == "387500140020"
    assert zurueck["EAN"].iloc[0] == "4004478123456"


def test_xlsx_haelt_lange_nummern_als_text(tmp_path, pool):
    ergebnis = matching.abgleich(pd.DataFrame({"Artikelnummer": ["387500140020"]}), pool)
    tabelle = export.importtabelle_bauen(ergebnis.importzeilen)
    dateien = export.schreibe_importdatei(tabelle, tmp_path, "Test")

    blatt = load_workbook(dateien["xlsx"])["Import"]
    kopf = [z.value for z in blatt[1]]
    for spalte in ("EAN", "Artikelnummer des Lieferanten"):
        zelle = blatt.cell(row=2, column=kopf.index(spalte) + 1)
        assert isinstance(zelle.value, str), f"{spalte} ist keine Zeichenkette"
        assert zelle.number_format == "@", f"{spalte} ist nicht als Text formatiert"
    assert blatt.cell(row=2, column=kopf.index("EAN") + 1).value == "4004478123456"


def test_xlsx_haelt_fuehrende_nullen_und_lange_nummern(tmp_path, pool):
    """Im echten Pool stehen 16.000+ rein numerische Artikelnummern ab 12 Stellen
    sowie Werte mit führenden Nullen. Excel würde sie ohne Textformat zerstören."""
    ergebnis = matching.abgleich(
        pd.DataFrame({"Artikelnummer": ["0000100951687", "02983"]}), pool
    )
    tabelle = export.importtabelle_bauen(ergebnis.importzeilen)
    dateien = export.schreibe_importdatei(tabelle, tmp_path, "Test")

    blatt = load_workbook(dateien["xlsx"])["Import"]
    kopf = [z.value for z in blatt[1]]
    for spalte in ("Artikelnummer des Lieferanten", "Werksnummer", "Herstellernummer"):
        for zeile in range(2, blatt.max_row + 1):
            zelle = blatt.cell(row=zeile, column=kopf.index(spalte) + 1)
            if zelle.value is None:
                continue
            assert isinstance(zelle.value, str), f"{spalte} wurde zur Zahl"
            assert zelle.number_format == "@", f"{spalte} ist nicht als Text formatiert"

    spalte_nr = kopf.index("Artikelnummer des Lieferanten") + 1
    gelesen = {blatt.cell(row=r, column=spalte_nr).value for r in (2, 3)}
    assert gelesen == {"0000100951687", "02983"}

    # auch in der CSV bleiben die führenden Nullen erhalten
    zurueck = pd.read_csv(dateien["csv"], sep=";", dtype=str, encoding="utf-8-sig")
    assert set(zurueck["Artikelnummer des Lieferanten"]) == {"0000100951687", "02983"}
    assert zurueck["Werksnummer"].dropna().tolist() == ["0015051500001"]


def test_xlsx_haelt_lager_min_max_als_zahl(tmp_path, pool):
    ergebnis = matching.abgleich(pd.DataFrame({"Artikelnummer": ["CUS18H"]}), pool)
    tabelle = export.importtabelle_bauen(ergebnis.importzeilen)
    dateien = export.schreibe_importdatei(tabelle, tmp_path, "Test")

    blatt = load_workbook(dateien["xlsx"])["Import"]
    kopf = [z.value for z in blatt[1]]
    assert blatt.cell(row=2, column=kopf.index("Lager Min") + 1).value == 4
    assert blatt.cell(row=2, column=kopf.index("Lager Max") + 1).value == 8


def test_report_wird_geschrieben(tmp_path, pool):
    ergebnis = matching.abgleich(pd.DataFrame({"Artikelnummer": ["CUS18H", "NEU-1"]}), pool)
    pfad = export.schreibe_report(pd.DataFrame(ergebnis.reportzeilen), tmp_path, "Test")

    blatt = load_workbook(pfad)["Report"]
    assert [z.value for z in blatt[1]][:2] == ["Artikelnummer", "Status"]
    assert blatt.max_row == 3


# ---------------------------------------------------------------------------
# Die Brücke zwischen Feldtabelle und Ausgabeformat
# ---------------------------------------------------------------------------
# Die Importzeile entsteht als dict aus FESTE_SPALTEN plus felder.FELDER.
# ``importtabelle_bauen`` sortiert anschließend auf die Vorlagenreihenfolge um –
# ein Feld, dessen ``importspalte`` es in der Vorlage nicht gibt, fiele dabei
# still unter den Tisch. Diese Tests halten beides fest.
def test_jedes_feld_zeigt_auf_eine_spalte_der_vorlage():
    from app import felder as felder_modul

    vorlage = set(export.vorlagenspalten())
    unbekannt = [f.name for f in felder_modul.FELDER if f.importspalte not in vorlage]
    assert unbekannt == [], f"Importspalte existiert nicht in der Vorlage: {unbekannt}"


def test_feste_spalten_und_felder_ergeben_zusammen_die_vorlage():
    from app import felder as felder_modul, matching as matching_modul

    erzeugt = set(matching_modul.FESTE_SPALTEN) | {
        f.importspalte for f in felder_modul.FELDER
    }
    assert erzeugt == set(export.vorlagenspalten())


def test_ein_echter_abgleich_besteht_die_formatpruefung(pool):
    """Der Weg, den die Anwendung wirklich geht – inklusive Umsortierung."""
    from app import matching as matching_modul

    ergebnis = matching_modul.abgleich(
        pd.DataFrame({"Artikelnummer": ["CUS28H", "CUS22H", "GIBTESNICHT-4711"]}), pool
    )
    df = export.importtabelle_bauen(ergebnis.importzeilen)
    assert list(df.columns) == export.vorlagenspalten()
    export.validiere_importformat(df)  # wirft, wenn etwas nicht stimmt
