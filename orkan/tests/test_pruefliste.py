"""Prüfliste: Export, Bearbeiten, Re-Import und die Qualitätsbremse."""
from __future__ import annotations

import io

import pandas as pd
import pytest
from openpyxl import load_workbook

from app import config, queue_store

ADMIN = ("pruefer", "geheim")


def _csv(inhalt: str) -> tuple[str, io.BytesIO, str]:
    return ("liste.csv", io.BytesIO(inhalt.encode("utf-8")), "text/csv")


def _fuelle_queue(klient, zeilen: str) -> list[str]:
    """Eine Kundenliste durchlaufen lassen und die entstandenen QueueIds liefern."""
    hoch = klient.post("/api/upload", files={"datei": _csv(zeilen)})
    assert hoch.status_code == 200, hoch.text
    job = hoch.json()
    antwort = klient.post("/api/verarbeiten", json={
        "job_id": job["job_id"], "zuordnung": job["zuordnung"],
        "kundenkennung": "Pruefkunde",
    })
    assert antwort.status_code == 200, antwort.text
    return list(queue_store.lesen(nur_offen=True)["QueueId"])


# ---------------------------------------------------------------------------
# Qualitätsprüfung
# ---------------------------------------------------------------------------
@pytest.mark.parametrize(
    "nummer,bezeichnung,erwarteter_grund",
    [
        ("Michael Seiler", "Bogen", "sieht nach Text aus"),
        ("Evoflex Gr. 8; Bund Orange", "Bogen", "Semikolon"),
        ("", "Bogen", "keine Artikelnummer"),
        ("123", "Bogen", "hoechstens vier Ziffern"),
        ("CUS99X", "", "keine Bezeichnung"),
        ("CUS99X", None, "keine Bezeichnung"),
    ],
)
def test_pruefung_erkennt_unbrauchbare_zeilen(nummer, bezeichnung, erwarteter_grund):
    gruende = queue_store.pruefe_eintrag(
        {"ArtikelNrLieferant": nummer, "BezArtikel1": bezeichnung}
    )
    assert any(erwarteter_grund in g for g in gruende), gruende


@pytest.mark.parametrize(
    "nummer",
    ["72 011 23", "93 022 89", "01.001.90", "0000100951687", "387500140020",
     "CUS18H", "QBSHNS", "QGCUHRA", "ALLARDSP"],
)
def test_pruefung_laesst_echte_pool_nummern_durch(nummer):
    """Über 5.000 gültige Pool-Nummern enthalten Leerzeichen – die dürfen nicht
    an der Prüfung scheitern."""
    assert queue_store.pruefe_eintrag(
        {"ArtikelNrLieferant": nummer, "BezArtikel1": "Irgendein Artikel"}
    ) == []


def test_pruefung_laesst_vollstaendige_zeile_durch():
    assert queue_store.pruefe_eintrag({
        "ArtikelNrLieferant": "CUS99X", "BezArtikel1": "Kupferbogen 90",
        "LagerMin": "2", "LagerMax": "8",
    }) == []


def test_pruefung_meckert_bei_min_groesser_max():
    gruende = queue_store.pruefe_eintrag({
        "ArtikelNrLieferant": "CUS99X", "BezArtikel1": "Bogen",
        "LagerMin": "9", "LagerMax": "3",
    })
    assert any("groesser" in g for g in gruende)


def test_pruefung_meckert_bei_unlesbarem_minmax():
    gruende = queue_store.pruefe_eintrag({
        "ArtikelNrLieferant": "CUS99X", "BezArtikel1": "Bogen",
        "LagerMin": "viele", "LagerMax": "3",
    })
    assert any("keine Zahl" in g for g in gruende)


# ---------------------------------------------------------------------------
# Übernahme wird blockiert
# ---------------------------------------------------------------------------
def test_uebernahme_lehnt_unvollstaendige_eintraege_ab(klient):
    """Genau der Fall, der 123 Zeilen ohne Bezeichnung in den Pool gespült hat."""
    ids = _fuelle_queue(klient, "Artikelnummer\nNEU-OHNE-BEZ-1\nNEU-OHNE-BEZ-2")
    vorher = len(pd.read_csv(config.POOL_PFAD, sep=";", dtype=str, encoding="utf-8-sig"))

    antwort = klient.post(
        "/admin/neuaufnahmen/uebernehmen", json={"queue_ids": ids}, auth=ADMIN
    )
    assert antwort.status_code == 400
    fehler = antwort.json()["fehler"]
    assert "keine Bezeichnung" in fehler
    assert "NEU-OHNE-BEZ-1" in fehler

    nachher = len(pd.read_csv(config.POOL_PFAD, sep=";", dtype=str, encoding="utf-8-sig"))
    assert nachher == vorher, "Der Pool darf dabei nicht angefasst werden"
    assert len(queue_store.lesen(nur_offen=True)) == len(ids)


def test_uebernahme_nennt_namen_statt_artikelnummer(klient):
    _fuelle_queue(klient, "Artikelnummer;Bezeichnung\nMichael Seiler;Irgendwas")
    ids = list(queue_store.lesen(nur_offen=True)["QueueId"])

    antwort = klient.post(
        "/admin/neuaufnahmen/uebernehmen", json={"queue_ids": ids}, auth=ADMIN
    )
    assert antwort.status_code == 400
    assert "sieht nach Text aus" in antwort.json()["fehler"]


def test_vollstaendiger_eintrag_wird_uebernommen(klient):
    _fuelle_queue(
        klient,
        "Artikelnummer;Bezeichnung;Artikelgruppe\nNEU-GUT-1;Kupferbogen 90;Pressfittings",
    )
    ids = list(queue_store.lesen(nur_offen=True)["QueueId"])

    antwort = klient.post(
        "/admin/neuaufnahmen/uebernehmen", json={"queue_ids": ids}, auth=ADMIN
    )
    assert antwort.status_code == 200
    assert antwort.json()["uebernommen"] == 1


# ---------------------------------------------------------------------------
# Export
# ---------------------------------------------------------------------------
def test_export_liefert_bearbeitbare_excel_datei(klient):
    _fuelle_queue(klient, "Artikelnummer\nNEU-EXP-1\nNEU-EXP-2")

    antwort = klient.get("/admin/neuaufnahmen/export", auth=ADMIN)
    assert antwort.status_code == 200
    assert "spreadsheetml" in antwort.headers["content-type"]

    blatt = load_workbook(io.BytesIO(antwort.content))["Neuaufnahmen"]
    kopf = [z.value for z in blatt[1]]
    assert kopf[:3] == ["Kennung", "Artikelnummer", "Bezeichnung 1"]
    assert queue_store.HINWEISSPALTE in kopf
    assert blatt.max_row - 1 == 2

    spalte = kopf.index(queue_store.HINWEISSPALTE) + 1
    assert "keine Bezeichnung" in blatt.cell(row=2, column=spalte).value


def test_export_braucht_anmeldung(klient):
    assert klient.get("/admin/neuaufnahmen/export").status_code == 401


def test_export_bei_leerer_queue_bleibt_gueltig(klient):
    antwort = klient.get("/admin/neuaufnahmen/export", auth=ADMIN)
    assert antwort.status_code == 200
    blatt = load_workbook(io.BytesIO(antwort.content))["Neuaufnahmen"]
    assert blatt.max_row == 1  # nur die Kopfzeile


# ---------------------------------------------------------------------------
# Re-Import
# ---------------------------------------------------------------------------
def _bearbeite_export(klient, aenderungen: dict, loeschen: set[str] = frozenset()):
    """Export holen, Zeilen ändern/löschen und als Excel zurückgeben."""
    roh = klient.get("/admin/neuaufnahmen/export", auth=ADMIN).content
    df = pd.read_excel(io.BytesIO(roh), dtype=object)
    for kennung, felder in aenderungen.items():
        for spalte, wert in felder.items():
            df.loc[df["Kennung"] == kennung, spalte] = wert
    if loeschen:
        df = df[~df["Kennung"].isin(loeschen)]
    puffer = io.BytesIO()
    df.to_excel(puffer, index=False)
    puffer.seek(0)
    return ("bearbeitet.xlsx", puffer,
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


def test_ergaenzte_bezeichnung_kommt_zurueck_und_erlaubt_uebernahme(klient):
    ids = _fuelle_queue(klient, "Artikelnummer\nNEU-IMP-1")
    kennung = ids[0]

    datei = _bearbeite_export(klient, {kennung: {
        "Bezeichnung 1": "Kupferbogen 90 Grad",
        "Artikelgruppe": "Pressfittings",
        "Lager Min": 3, "Lager Max": 9,
    }})
    antwort = klient.post("/admin/neuaufnahmen/import", files={"datei": datei}, auth=ADMIN)
    assert antwort.status_code == 200, antwort.text
    assert antwort.json()["aktualisiert"] == 1
    assert antwort.json()["unvollstaendig"] == 0

    eintrag = queue_store.lesen(nur_offen=True).iloc[0]
    assert eintrag["BezArtikel1"] == "Kupferbogen 90 Grad"
    assert eintrag["BezArtikelGruppe"] == "Pressfittings"

    uebernahme = klient.post(
        "/admin/neuaufnahmen/uebernehmen", json={"queue_ids": [kennung]}, auth=ADMIN
    )
    assert uebernahme.status_code == 200
    assert uebernahme.json()["uebernommen"] == 1


def test_geloeschte_zeile_gilt_als_verworfen(klient):
    ids = _fuelle_queue(klient, "Artikelnummer\nNEU-DEL-1\nNEU-DEL-2")
    behalten, weg = ids[0], ids[1]

    datei = _bearbeite_export(
        klient, {behalten: {"Bezeichnung 1": "Bleibt drin"}}, loeschen={weg}
    )
    antwort = klient.post("/admin/neuaufnahmen/import", files={"datei": datei}, auth=ADMIN)
    assert antwort.status_code == 200
    assert antwort.json() == {
        "aktualisiert": 1, "verworfen": 1, "unbekannt": [],
        "offen": 1, "unvollstaendig": 0,
    }

    alle = queue_store.lesen()
    assert alle.loc[alle["QueueId"] == weg, "QueueStatus"].iloc[0] == queue_store.STATUS_VERWORFEN
    assert alle.loc[alle["QueueId"] == behalten, "QueueStatus"].iloc[0] == queue_store.STATUS_OFFEN


def test_artikelnummer_laesst_sich_nicht_ueber_den_import_aendern(klient):
    ids = _fuelle_queue(klient, "Artikelnummer\nNEU-FEST-1")
    datei = _bearbeite_export(klient, {ids[0]: {
        "Artikelnummer": "GANZ-ANDERE-NUMMER", "Bezeichnung 1": "Bogen",
    }})
    klient.post("/admin/neuaufnahmen/import", files={"datei": datei}, auth=ADMIN)

    assert queue_store.lesen(nur_offen=True).iloc[0]["ArtikelNrLieferant"] == "NEU-FEST-1"


def test_import_ohne_kennungsspalte_meldet_klaren_fehler(klient):
    _fuelle_queue(klient, "Artikelnummer\nNEU-X-1")
    puffer = io.BytesIO()
    pd.DataFrame({"Irgendwas": ["abc"]}).to_excel(puffer, index=False)
    puffer.seek(0)

    antwort = klient.post(
        "/admin/neuaufnahmen/import",
        files={"datei": ("falsch.xlsx", puffer, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        auth=ADMIN,
    )
    assert antwort.status_code == 400
    assert "Kennung" in antwort.json()["fehler"]


def test_import_meldet_unbekannte_kennungen(klient):
    ids = _fuelle_queue(klient, "Artikelnummer\nNEU-U-1")
    roh = klient.get("/admin/neuaufnahmen/export", auth=ADMIN).content
    df = pd.read_excel(io.BytesIO(roh), dtype=object)
    df.loc[len(df)] = {"Kennung": "gibtesnicht-0001", "Bezeichnung 1": "Fantasie"}
    puffer = io.BytesIO()
    df.to_excel(puffer, index=False)
    puffer.seek(0)

    antwort = klient.post(
        "/admin/neuaufnahmen/import",
        files={"datei": ("b.xlsx", puffer, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        auth=ADMIN,
    )
    assert antwort.status_code == 200
    assert antwort.json()["unbekannt"] == ["gibtesnicht-0001"]
    assert len(queue_store.lesen(nur_offen=True)) == len(ids)


def test_import_braucht_anmeldung(klient):
    puffer = io.BytesIO(b"x")
    antwort = klient.post(
        "/admin/neuaufnahmen/import", files={"datei": ("a.xlsx", puffer, "application/octet-stream")}
    )
    assert antwort.status_code == 401


def test_import_lehnt_fremde_dateiformate_ab(klient):
    antwort = klient.post(
        "/admin/neuaufnahmen/import",
        files={"datei": ("bild.png", io.BytesIO(b"\x89PNG"), "image/png")},
        auth=ADMIN,
    )
    assert antwort.status_code == 400
    assert "Excel" in antwort.json()["fehler"]


# ---------------------------------------------------------------------------
# Oberfläche
# ---------------------------------------------------------------------------
def test_admin_seite_zeigt_pruefhinweise_und_bedienelemente(klient):
    _fuelle_queue(klient, "Artikelnummer\nNEU-UI-1")
    seite = klient.get("/admin/neuaufnahmen", auth=ADMIN).text

    assert "Als Excel herunterladen" in seite
    assert "Einspielen" in seite
    assert "In Pool übernehmen" in seite
    assert "keine Bezeichnung" in seite
    assert "unvollständig" in seite


# ---------------------------------------------------------------------------
# Erreichbarkeit je Betriebsart
# ---------------------------------------------------------------------------
def test_status_meldet_die_pruefliste(klient):
    daten = klient.get("/api/status").json()["pruefliste"]
    assert daten["erreichbar"] is True          # Zugangsdaten sind in den Tests gesetzt
    assert daten["pool_beschreibbar"] is True
    assert isinstance(daten["offen"], int)


def test_desktop_erreicht_die_pruefliste_ohne_anmeldung(klient, monkeypatch):
    """In der .exe lauscht der Dienst nur lokal – eine Anmeldung braucht es nicht."""
    monkeypatch.setattr(config, "gebuendelt", lambda: True)
    antwort = klient.get("/admin/neuaufnahmen")          # ohne auth=
    assert antwort.status_code == 200


def test_desktop_zeigt_keinen_uebernehmen_knopf(klient, monkeypatch):
    monkeypatch.setattr(config, "gebuendelt", lambda: True)
    monkeypatch.setattr(config, "pool_schreibbar", lambda: False)
    seite = klient.get("/admin/neuaufnahmen").text

    assert "In Pool übernehmen" not in seite
    assert "an OrKan übergeben" in seite
    assert "Als Excel herunterladen" in seite      # Export bleibt möglich
    assert "Verwerfen" in seite


def test_desktop_kann_die_liste_weiterhin_exportieren(klient, monkeypatch):
    _fuelle_queue(klient, "Artikelnummer;Bezeichnung\nNEU-DESK-1;Ein Artikel")
    monkeypatch.setattr(config, "gebuendelt", lambda: True)

    antwort = klient.get("/admin/neuaufnahmen/export")   # ohne Anmeldung
    assert antwort.status_code == 200
    assert "spreadsheetml" in antwort.headers["content-type"]


def test_server_verlangt_weiterhin_eine_anmeldung(klient, monkeypatch):
    monkeypatch.setattr(config, "gebuendelt", lambda: False)
    assert klient.get("/admin/neuaufnahmen").status_code == 401
