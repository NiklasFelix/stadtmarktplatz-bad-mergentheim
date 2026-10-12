"""Artikelnummern müssen zeichengenau erhalten bleiben – über die ganze Kette.

Ausdrückliche Vorgabe: Nummern wie ``72 011 23`` sind gültig und werden
akzeptiert. Im Pool stehen davon 5.635 Stück. Sie dürfen weder abgelehnt noch
umgeschrieben werden – nicht beim Einlesen, nicht beim Abgleich, nicht in der
Ausgabedatei und nicht in der Prüf-Queue.

Ebenso gültig: Nummern ohne jede Ziffer (``QBSHNS``, 1.240 Stück), mit
führenden Nullen (``0000100951687``, 919 Stück) und mit Punkten (``01.001.90``).
"""
from __future__ import annotations

import io

import pandas as pd
import pytest
from openpyxl import load_workbook

from app import export, matching, parsing, queue_store

#: Schreibweisen, die im echten Pool vorkommen und alle gültig sind.
ECHTE_SCHREIBWEISEN = [
    "72 011 23",       # Leerzeichen als Gruppierung – 5.635 im Pool
    "93 022 89",
    "QBSHNS",          # gar keine Ziffer – 1.240 im Pool
    "QGCUHRA",
    "01.001.90",       # Punkte
    "0000100951687",   # führende Nullen – 919 im Pool
    "387500140020",    # lange reine Zahl – 16.064 im Pool
    "CUS18H",
]


def _csv(inhalt: str):
    return ("liste.csv", io.BytesIO(inhalt.encode("utf-8")), "text/csv")


# ---------------------------------------------------------------------------
# Einlesen
# ---------------------------------------------------------------------------
@pytest.mark.parametrize("nummer", ECHTE_SCHREIBWEISEN)
def test_einlesen_veraendert_die_nummer_nicht(nummer):
    roh = pd.DataFrame({"Artikelnummer": [nummer], "Bezeichnung": ["Testartikel"]})
    liste = parsing.baue_kundenliste(
        roh, {"Artikelnummer": "Artikelnummer", "Bezeichnung1": "Bezeichnung"}
    )
    assert liste["Artikelnummer"].iloc[0] == nummer


@pytest.mark.parametrize("nummer", ECHTE_SCHREIBWEISEN)
def test_normalisierung_behaelt_innere_leerzeichen(nummer):
    """Getrimmt und großgeschrieben wird – innen bleibt alles stehen."""
    assert matching.normiere_artikelnummer(f"  {nummer}  ") == nummer.upper()


def test_spaltenerkennung_findet_nummern_mit_leerzeichen():
    """Ohne passenden Header muss die Heuristik solche Nummern erkennen."""
    roh = pd.DataFrame({
        "Spalte A": ["Bogen 90 Grad", "Muffe rund", "Rohr lang"],
        "Spalte B": ["72 011 23", "93 022 89", "97 009 32"],
    })
    assert parsing.erkenne_spalten(roh).artikelnummer_spalte == "Spalte B"


# ---------------------------------------------------------------------------
# Abgleich
# ---------------------------------------------------------------------------
def test_nummern_mit_leerzeichen_gelten_nicht_als_ungueltig():
    for nummer in ECHTE_SCHREIBWEISEN:
        assert not matching.ist_ungueltige_nummer(nummer), nummer


@pytest.mark.parametrize("nummer", ["72 011 23", "QBSHNS"])
def test_solche_nummern_werden_im_pool_gefunden(pool, nummer):
    """Der Treffer-Weg: die Nummer muss den Pool-Artikel finden."""
    ergebnis = matching.abgleich(pd.DataFrame({"Artikelnummer": [nummer]}), pool)

    assert ergebnis.anzahl_gefunden == 1
    assert ergebnis.anzahl_ungueltig == 0
    assert ergebnis.importzeilen[0]["Artikelnummer des Lieferanten"] == nummer
    assert ergebnis.importzeilen[0]["Bezeichnung1"] in (
        "Waschtischarmatur", "Schlauchbrause"
    )


@pytest.mark.parametrize("schreibweise", ["72 011 23", " 72 011 23 ", "72 011 23\n"])
def test_treffer_auch_bei_umliegenden_leerzeichen(pool, schreibweise):
    ergebnis = matching.abgleich(pd.DataFrame({"Artikelnummer": [schreibweise]}), pool)
    assert ergebnis.anzahl_gefunden == 1


def test_abgleich_reicht_die_nummer_unveraendert_durch(pool):
    """Auch als Neuaufnahme bleibt die Schreibweise erhalten."""
    liste = pd.DataFrame([{"Artikelnummer": "88 999 77", "Bezeichnung1": "Neuer Artikel"}])
    ergebnis = matching.abgleich(liste, pool)

    assert ergebnis.anzahl_neu == 1
    assert ergebnis.neuaufnahmen[0]["ArtikelNrLieferant"] == "88 999 77"
    assert ergebnis.importzeilen[0]["Artikelnummer des Lieferanten"] == "88 999 77"
    assert ergebnis.reportzeilen[0]["Artikelnummer"] == "88 999 77"


# ---------------------------------------------------------------------------
# Qualitätsbremse
# ---------------------------------------------------------------------------
@pytest.mark.parametrize("nummer", ECHTE_SCHREIBWEISEN)
def test_qualitaetsbremse_akzeptiert_alle_echten_schreibweisen(nummer):
    assert queue_store.pruefe_eintrag(
        {"ArtikelNrLieferant": nummer, "BezArtikel1": "Testartikel"}
    ) == []


@pytest.mark.parametrize(
    "nummer", ["Michael Seiler", "Andreas Funkner", "Abflussreiniger Buzil"]
)
def test_qualitaetsbremse_faengt_weiterhin_echte_namen_ab(nummer):
    """Die Lockerung darf Namen nicht durchlassen."""
    gruende = queue_store.pruefe_eintrag(
        {"ArtikelNrLieferant": nummer, "BezArtikel1": "Testartikel"}
    )
    assert any("sieht nach Text aus" in g for g in gruende), gruende


# ---------------------------------------------------------------------------
# Ausgabedateien
# ---------------------------------------------------------------------------
def test_ausgabedateien_halten_die_schreibweise(tmp_path, pool):
    liste = pd.DataFrame({
        "Artikelnummer": ECHTE_SCHREIBWEISEN,
        "Bezeichnung1": ["Testartikel"] * len(ECHTE_SCHREIBWEISEN),
    })
    ergebnis = matching.abgleich(liste, pool)
    tabelle = export.importtabelle_bauen(ergebnis.importzeilen)
    dateien = export.schreibe_importdatei(tabelle, tmp_path, "Formate")

    # CSV
    zurueck = pd.read_csv(dateien["csv"], sep=";", dtype=str, encoding="utf-8-sig")
    assert zurueck["Artikelnummer des Lieferanten"].tolist() == ECHTE_SCHREIBWEISEN

    # Excel: als Text formatiert, damit Excel nichts umrechnet
    blatt = load_workbook(dateien["xlsx"])["Import"]
    kopf = [z.value for z in blatt[1]]
    spalte = kopf.index("Artikelnummer des Lieferanten") + 1
    gelesen = []
    for zeile in range(2, blatt.max_row + 1):
        zelle = blatt.cell(row=zeile, column=spalte)
        assert isinstance(zelle.value, str), f"{zelle.value!r} wurde zur Zahl"
        assert zelle.number_format == "@"
        gelesen.append(zelle.value)
    assert gelesen == ECHTE_SCHREIBWEISEN


# ---------------------------------------------------------------------------
# Kompletter Weg über die Weboberfläche
# ---------------------------------------------------------------------------
def test_ganzer_ablauf_haelt_die_schreibweise(klient):
    inhalt = "Artikelnummer;Bezeichnung\n" + "\n".join(
        f"{n};Testartikel" for n in ECHTE_SCHREIBWEISEN
    )
    hoch = klient.post("/api/upload", files={"datei": _csv(inhalt)})
    assert hoch.status_code == 200
    job = hoch.json()

    antwort = klient.post("/api/verarbeiten", json={
        "job_id": job["job_id"], "zuordnung": job["zuordnung"],
        "kundenkennung": "Formattest",
    })
    assert antwort.status_code == 200
    ergebnis = antwort.json()

    assert ergebnis["zusammenfassung"]["ungueltig"] == 0
    assert [z["Artikelnummer"] for z in ergebnis["report"]] == ECHTE_SCHREIBWEISEN

    datei = klient.get(ergebnis["downloads"]["csv"])
    tabelle = pd.read_csv(
        io.BytesIO(datei.content), sep=";", dtype=str, encoding="utf-8-sig"
    )
    assert tabelle["Artikelnummer des Lieferanten"].tolist() == ECHTE_SCHREIBWEISEN


def test_pruefliste_zeigt_keine_beanstandung_fuer_solche_nummern(klient):
    inhalt = "Artikelnummer;Bezeichnung\n72 011 23;Vollstaendiger Artikel"
    hoch = klient.post("/api/upload", files={"datei": _csv(inhalt)})
    klient.post("/api/verarbeiten", json={
        "job_id": hoch.json()["job_id"], "zuordnung": hoch.json()["zuordnung"],
        "kundenkennung": "Formattest",
    })

    offen = queue_store.lesen(nur_offen=True)
    if len(offen):  # nur wenn die Nummer im Testpool fehlt
        assert queue_store.pruefe_eintrag(offen.iloc[0]) == []
