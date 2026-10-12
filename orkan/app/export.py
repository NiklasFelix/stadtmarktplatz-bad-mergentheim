"""Erzeugen und Validieren der Ausgabedateien.

Die 39 Spalten der Importdatei kommen aus der Vorlage
``referenz/Importformat_39_Spalten_Vorlage.csv``. Vor jedem Schreibvorgang wird
geprüft, dass Spaltennamen und Reihenfolge exakt übereinstimmen – lieber ein
Fehler als eine falsche Datei.
"""
from __future__ import annotations

import csv
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from . import config

#: Spalten, die Excel sonst als Zahl interpretiert und dabei zerstört
#: (lange EANs werden zu 4,04478E+12, führende Nullen fallen weg).
TEXTSPALTEN = {
    "Artikelnr.",
    "Artikelnummer des Lieferanten",
    "EAN",
    "Werksnummer",
    "Herstellernummer",
    "HTN-Code",
    "Matchcode",
    "Artikelnummer",
}

#: Spalten, die als echte Zahl in die Tabelle sollen.
ZAHLSPALTEN = {"Lager Min", "Lager Max", "Menge(VPE)", "Bestand", "Abw. Bestellmenge"}

_vorlage_cache: list[str] | None = None


def vorlagenspalten(pfad: Path | None = None) -> list[str]:
    """Die 39 verbindlichen Spaltennamen in verbindlicher Reihenfolge."""
    global _vorlage_cache
    if pfad is None and _vorlage_cache is not None:
        return list(_vorlage_cache)

    ziel = Path(pfad or config.VORLAGE_PFAD)
    with ziel.open("r", encoding=config.CSV_KODIERUNG, newline="") as f:
        kopfzeile = next(csv.reader(f, delimiter=config.CSV_TRENNZEICHEN))
    spalten = [s.strip() for s in kopfzeile]
    if not spalten:
        raise ValueError(f"Formatvorlage {ziel} ist leer.")
    if pfad is None:
        _vorlage_cache = list(spalten)
    return spalten


def validiere_importformat(df: pd.DataFrame) -> None:
    """Spaltennamen und Reihenfolge exakt gegen die Vorlage prüfen."""
    erwartet = vorlagenspalten()
    tatsaechlich = list(df.columns)
    if tatsaechlich == erwartet:
        return

    fehlend = [s for s in erwartet if s not in tatsaechlich]
    zuviel = [s for s in tatsaechlich if s not in erwartet]
    details = []
    if fehlend:
        details.append(f"fehlende Spalten: {', '.join(fehlend)}")
    if zuviel:
        details.append(f"unbekannte Spalten: {', '.join(zuviel)}")
    if not details:
        details.append("Reihenfolge weicht von der Vorlage ab")
    raise ValueError(
        "Die Importdatei entspricht nicht dem vorgegebenen 39-Spalten-Format "
        f"({len(erwartet)} Spalten erwartet, {len(tatsaechlich)} erhalten): "
        + "; ".join(details)
    )


def leere_importtabelle() -> pd.DataFrame:
    return pd.DataFrame(columns=vorlagenspalten())


def importtabelle_bauen(zeilen: list[dict]) -> pd.DataFrame:
    """Importzeilen in eine Tabelle mit exakt den Vorlagenspalten gießen."""
    spalten = vorlagenspalten()
    df = pd.DataFrame(zeilen, columns=spalten) if zeilen else leere_importtabelle()
    df = df.reindex(columns=spalten)
    validiere_importformat(df)
    return df


# ---------------------------------------------------------------------------
# Schreiben
# ---------------------------------------------------------------------------
def _als_text(wert) -> str:
    if wert is None:
        return ""
    if isinstance(wert, float) and pd.isna(wert):
        return ""
    return str(wert)


def _als_zahl(wert) -> int | float | None:
    """Zahlenwert für Excel; leer wenn nicht als Zahl lesbar."""
    zahl = pd.to_numeric(wert, errors="coerce")
    if pd.isna(zahl):
        return None
    # pandas liefert je nach Eingabe numpy-Skalare oder normale Python-Zahlen.
    zahl = zahl.item() if hasattr(zahl, "item") else zahl
    return int(zahl) if float(zahl).is_integer() else float(zahl)


def schreibe_csv(df: pd.DataFrame, ziel: Path) -> Path:
    """Semikolon-getrennt, UTF-8 mit BOM – öffnet in Excel ohne kaputte Umlaute."""
    ziel.parent.mkdir(parents=True, exist_ok=True)
    df.to_csv(
        ziel,
        sep=config.CSV_TRENNZEICHEN,
        index=False,
        encoding=config.CSV_KODIERUNG,
        na_rep="",
        quoting=csv.QUOTE_MINIMAL,
        lineterminator="\r\n",
    )
    return ziel


def schreibe_xlsx(df: pd.DataFrame, ziel: Path, blattname: str = "Daten") -> Path:
    """Excel-Datei schreiben; Nummernspalten bleiben garantiert Text."""
    ziel.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = blattname[:31]

    spalten = list(df.columns)
    ws.append(spalten)
    for zelle in ws[1]:
        zelle.font = Font(bold=True)

    for _, zeile in df.iterrows():
        werte = []
        for spalte in spalten:
            roh = zeile[spalte]
            if spalte in ZAHLSPALTEN:
                werte.append(_als_zahl(roh))
            else:
                text = _als_text(roh)
                werte.append(None if text == "" else text)
        ws.append(werte)

    # Nummernspalten explizit als Text formatieren, damit Excel lange EANs
    # nicht in wissenschaftliche Notation umwandelt.
    for i, spalte in enumerate(spalten, start=1):
        buchstabe = get_column_letter(i)
        breite = max(10, min(38, len(str(spalte)) + 4))
        ws.column_dimensions[buchstabe].width = breite
        if spalte in TEXTSPALTEN:
            for zelle in ws[buchstabe]:
                if zelle.row > 1:
                    zelle.number_format = "@"

    ws.freeze_panes = "A2"
    wb.save(ziel)
    return ziel


def schreibe_importdatei(df: pd.DataFrame, verzeichnis: Path, name: str) -> dict[str, Path]:
    """Importdatei als .xlsx und .csv – nach erneuter Formatprüfung."""
    validiere_importformat(df)
    verzeichnis.mkdir(parents=True, exist_ok=True)
    return {
        "xlsx": schreibe_xlsx(df, verzeichnis / f"Importdatei_{name}.xlsx", "Import"),
        "csv": schreibe_csv(df, verzeichnis / f"Importdatei_{name}.csv"),
    }


def schreibe_report(df: pd.DataFrame, verzeichnis: Path, name: str) -> Path:
    verzeichnis.mkdir(parents=True, exist_ok=True)
    return schreibe_xlsx(df, verzeichnis / f"Abgleich_Report_{name}.xlsx", "Report")
